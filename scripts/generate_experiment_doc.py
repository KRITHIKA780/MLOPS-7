from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "airflow_retraining_experiment_guide.docx"


def add_code(document, text):
    paragraph = document.add_paragraph()
    paragraph.style = "No Spacing"
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

document.add_heading("Airflow Automated Model Retraining Experiment", 0)
document.add_paragraph(
    "Complete procedure for simulating performance degradation, detecting it, "
    "retraining an Iris classifier, replacing the production model, and sending alerts."
)

document.add_heading("1. Prerequisites", 1)
document.add_paragraph(
    "Windows 10/11, Docker Desktop with the Linux engine running, Docker Compose, "
    "and this project opened in VS Code. The Python scripts can also be run locally "
    "with Python 3.10 or newer and the packages in requirements.txt."
)

document.add_heading("2. Project Setup", 1)
document.add_paragraph("Open PowerShell and run:")
add_code(document, "cd D:\\EXP-7-airflow-retraining\n"
                   "python -m pip install -r requirements.txt\n"
                   "python scripts/train_model.py")
document.add_paragraph(
    "The training command creates models/production_model.pkl and data/iris_test.csv. "
    "The Docker image installs the same ML dependencies through the project Dockerfile."
)

document.add_heading("3. Simulate Degradation Locally", 1)
add_code(document, "python scripts/degrade_data.py\n"
                   "python scripts/evaluate_model.py\n"
                   "python scripts/check_degradation.py")
document.add_paragraph(
    "The experiment adds Gaussian noise to the Iris test features. The expected degraded "
    "accuracy is below the 0.80 threshold; the observed value in this run was 0.5667. "
    "check_degradation.py exits successfully when retraining is required."
)

document.add_heading("4. Start Airflow", 1)
add_code(document, "docker desktop start\n"
                   "docker compose up -d --build\n"
                   "docker compose ps")
document.add_paragraph(
    "Open http://localhost:8080. Airflow standalone prints the generated admin password "
    "in its logs. Retrieve it with the following command if needed:")
add_code(document, "docker compose logs airflow | Select-String -Pattern \"Password\"")

document.add_heading("5. Configure Optional Notifications", 1)
document.add_paragraph(
    "Every completion or failure is always written to data/notifications.log. To also post "
    "a JSON message to a compatible webhook, set the variable before starting the stack:")
add_code(document, "$env:NOTIFICATION_WEBHOOK_URL = \"https://your-webhook-endpoint\"\n"
                   "docker compose up -d --build")
document.add_paragraph(
    "The DAG sends a success notification after deployment and a failure notification when "
    "a task fails. Leave the variable empty to use the local audit log only."
)

document.add_heading("6. Trigger the Automated DAG", 1)
add_code(document, "docker compose exec airflow airflow dags list\n"
                   "docker compose exec airflow airflow dags trigger automatic_model_retraining\n"
                   "docker compose exec airflow airflow dags list-runs automatic_model_retraining -o plain")
document.add_paragraph(
    "The task order is simulate_degradation -> evaluate_model -> check_degradation. If the "
    "accuracy is below 0.80, the BranchPythonOperator selects retrain_model -> deploy_model. "
    "Otherwise it selects model_ok. Both paths finish with a notification task."
)

document.add_heading("7. Verify the Result", 1)
add_code(document, "docker compose exec airflow cat /opt/airflow/data/current_accuracy.txt\n"
                   "docker compose exec airflow cat /opt/airflow/data/new_accuracy.txt\n"
                   "docker compose exec airflow cat /opt/airflow/data/notifications.log\n"
                   "docker compose exec airflow ls -l /opt/airflow/models")
document.add_paragraph(
    "The retraining task evaluates a fresh RandomForest model on a clean holdout split and "
    "writes new_accuracy.txt. deploy_model.py copies new_model.pkl over production_model.pkl. "
    "In the local verification, new_accuracy.txt was 0.90 and the replacement completed."
)

document.add_heading("8. Useful Operations", 1)
add_code(document, "docker compose logs -f airflow\n"
                   "docker compose down\n"
                   "docker compose down -v  # also removes the Postgres volume\n"
                   "docker compose up -d --build")

document.add_heading("9. Troubleshooting", 1)
document.add_paragraph(
    "If Compose reports that it cannot connect to dockerDesktopLinuxEngine, start Docker "
    "Desktop and wait until its engine is ready, then rerun docker compose up -d --build. "
    "If a DAG does not appear, inspect docker compose logs airflow and confirm that the DAG "
    "file is mounted at /opt/airflow/dags."
)

document.add_heading("10. Expected Deliverables", 1)
document.add_paragraph(
    "After a degraded run, the workspace contains data/degraded_data.csv, current_accuracy.txt, "
    "new_accuracy.txt, notifications.log, models/new_model.pkl, and the replaced "
    "models/production_model.pkl. The Airflow UI shows the automatic_model_retraining DAG "
    "and its task graph."
)

document.save(OUTPUT)
print(OUTPUT)